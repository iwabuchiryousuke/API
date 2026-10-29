"""Actual spreadsheet recalculation, using synthetic records and LibreOffice.

Office is optional for downstream developer machines; no Office package is an
application dependency. This suite never launches SBI, Gmail, or Windows tasks.
"""
from datetime import date,datetime,timedelta,timezone
import hashlib
import json
from pathlib import Path
import os
import shutil
import subprocess
import tempfile
import unittest

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook=None

from activity import Activities
from current_capture import save_current_capture
from ledger import Snapshot,Store
from report_v36 import build
from standalone_excel import write_model
from split_entry import register


JST=timezone(timedelta(hours=9))
OFFICE_FALLBACK=Path('/opt/codex/runtimes/codex-primary-runtime/dependencies/bin/override/soffice')


def office_executable():
    requested=os.environ.get('SBI_TEST_OFFICE')
    if requested and Path(requested).is_file():return str(Path(requested))
    found=shutil.which('soffice') or shutil.which('libreoffice')
    if found:return found
    return str(OFFICE_FALLBACK) if OFFICE_FALLBACK.is_file() else None


def create_office_fixture(base):
    """Use prices chosen to reveal once-per-day versus per-fill ceiling."""
    base=Path(base);store=Store(base/'data');activities=Activities(store)
    for identifier in 'ABCDE':
        store.rename(identifier,'Account '+identifier)
        store.set_owner(identifier,'Example '+identifier)
    def trade(day,quantity,cash,side='買',code='9900',name='Example Security'):
        return dict(trade_date=day,settlement_date=day,name=name,code=code,market='東証',
                    side=side,custody='NISA(成)',tax_class='',quantity=str(quantity),
                    price=str(cash/quantity),fees='0',tax='0',settlement_amount=str(cash),
                    asset_type='国内株式',quantity_unit='株',price_unit='円/株',currency='JPY')
    records=[trade('2026-10-01',3,301),trade('2026-10-01',1,103),
             trade('2026-10-02',2,240,'売'),trade('2026-10-03',2,205),
             trade('2026-10-05',3,180,'売'),
             trade('2026-10-06',1,200,'売','9901','Missing History Security')]
    raw=json.dumps(records,sort_keys=True).encode()
    store.import_snapshot('A',Snapshot(raw,'office-synthetic.csv',hashlib.sha256(raw).hexdigest(),
                                      '2026-10-01','2026-10-06',records))
    ordered=[trade('2026-10-01',1,100,code='9902',name='Daily Ordering Security'),
             trade('2026-10-02',3,301,code='9902',name='Daily Ordering Security'),
             trade('2026-10-02',1,120,'売','9902','Daily Ordering Security')]
    raw=json.dumps(ordered,sort_keys=True).encode()
    store.import_snapshot('B',Snapshot(raw,'office-ordering.csv',hashlib.sha256(raw).hexdigest(),
                                      '2026-10-01','2026-10-02',ordered))
    funds=[]
    for day,quantity,cash,side in (('2026-10-01',10000,10001,'買'),
                                  ('2026-10-01',10000,10003,'買'),
                                  ('2026-10-02',10000,10500,'売')):
        item=trade(day,quantity,cash,side,'9903','Fund Quote Unit Security')
        item.update(asset_type='投資信託',quantity_unit='口',price_unit='円/1万口',price=str(cash/quantity*10000))
        funds.append(item)
    raw=json.dumps(funds,sort_keys=True).encode()
    store.import_snapshot('C',Snapshot(raw,'office-fund.csv',hashlib.sha256(raw).hexdigest(),
                                      '2026-10-01','2026-10-02',funds))
    register(store,activities,'A','9900','2026-10-04',1,2,'Synthetic split fixture')
    activities.add_cash('A','配当','2026-10-15','15','Example Security',
                        code='9900',custody='NISA(成)',quantity='3',per_share='5')
    activities.add_cash('A','税還付','2026-10-20','20','Tax refund')
    activities.add_cash('E','利息','2026-10-21','5','Interest')
    activities.add_cash('D','税還付','2025-10-20','7','Previous tax refund')
    activities.add_cash('D','利息','2025-10-21','2','Previous interest')
    market=store.directory/'market';market.mkdir()
    fact=dict(calendar_complete=True,rights='2026/10/27',payments='2026/11/10',
              earnings='2026/11/11',dps=4,close=150,close_date='2026-09-30',
              month_start_close=155,month_start_close_date='2026-10-01',
              events=[dict(rights='2026-10-27',payment='2026-11-10',dps=4)])
    (market/'research_2026.json').write_text(json.dumps({'year':2026,'securities':{'9900':fact}}))
    capture={}
    for index,identifier in enumerate('ABCD'):
        stamp=datetime(2026,10,29,15,10+index,tzinfo=JST).isoformat()
        values=dict(cash_total=0 if identifier=='B' else 10000+index,
                    growth_remaining={'year':2026,'remaining':1000+index},
                    tsumitate_remaining={'year':2026,'remaining':2000+index})
        sources={'口座サマリー':stamp}
        if identifier!='D':
            values['usd_cash']=0 if identifier=='B' else 123.45+index
            sources['口座(外貨建)']=stamp
        capture[identifier]=dict(owner='Example '+identifier,observed_at=stamp,values=values,
                                missing=[] if identifier!='D' else ['usd_cash'],source_observed_at=sources)
    save_current_capture(store,'office-current-fixture',capture,'2026-10-29T15:00:00+09:00')
    return store,activities


class ActualOfficeWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if load_workbook is None:raise unittest.SkipTest('openpyxl is needed only for optional Office validation')
        executable=office_executable()
        if executable is None:raise unittest.SkipTest('LibreOffice is not installed; set SBI_TEST_OFFICE to its executable')
        temporary=tempfile.TemporaryDirectory(prefix='sbi_v36_office_')
        cls.addClassCleanup(temporary.cleanup);cls.directory=Path(temporary.name)
        store,activities=create_office_fixture(cls.directory)
        cls.addClassCleanup(store.close)
        before={table:sorted(tuple(row) for row in store.db.execute('SELECT * FROM '+table))
                for table in ('accounts','imports','trades','evidence','activities','current_acquisitions')}
        cls.model=build(store,activities,'2026-10',today=date(2026,10,29))
        if any(sheet.get('freeze')!='C2' or sheet.get('header_height',0)<38 for sheet in cls.model['sheets']):
            raise AssertionError('Production readable-layout styling was not applied')
        cls.original=write_model(cls.model,cls.directory/'source')
        after={table:sorted(tuple(row) for row in store.db.execute('SELECT * FROM '+table)) for table in before}
        if before!=after:raise AssertionError('Workbook construction modified the synthetic ledger')
        converted=cls.directory/'recalculated';converted.mkdir()
        command=[executable,'-env:UserInstallation='+(cls.directory/'office-profile').as_uri(),
                 '--headless','--convert-to','xlsx','--outdir',str(converted),str(cls.original)]
        result=subprocess.run(command,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,timeout=90)
        cls.calculated=converted/cls.original.name
        if result.returncode!=0 or not cls.calculated.is_file():
            raise AssertionError('LibreOffice did not produce a recalculated synthetic workbook: '+result.stdout+result.stderr)
        cls.values=load_workbook(cls.calculated,data_only=True)
        cls.formulas=load_workbook(cls.calculated,data_only=False)
        cls.addClassCleanup(cls.values.close);cls.addClassCleanup(cls.formulas.close)

    def test_running_cost_formulas_recalculate_daily_group_split_and_sales(self):
        sheet=self.values['取引'];formulas=self.formulas['取引']
        found={}
        for row in range(2,sheet.max_row+1):
            if not isinstance(sheet.cell(row,1).value,datetime):continue
            key=(sheet.cell(row,1).value.date().isoformat(),str(sheet.cell(row,4).value))
            found[key]=(sheet.cell(row,7).value,sheet.cell(row,9).value,
                        sheet.cell(row,11).value,sheet.cell(row,12).value,sheet.cell(row,13).value)
            for column in (11,12,13):self.assertEqual(formulas.cell(row,column).data_type,'f')
        self.assertEqual(found[('2026-10-01','9900')],(4,404,4,101,0))
        self.assertEqual(found[('2026-10-02','9900')],(2,240,2,101,38))
        self.assertEqual(found[('2026-10-03','9900')],(2,205,4,102,0))
        self.assertEqual(found[('2026-10-04','9900')],(4,0,8,51,0))
        self.assertEqual(found[('2026-10-05','9900')],(3,180,5,51,27))
        self.assertEqual(len([key for key in found if key==('2026-10-01','9900')]),1)

    def test_incomplete_history_remains_blank_and_no_formula_has_an_error(self):
        sheet=self.values['取引'];unknown=None
        for row in range(2,sheet.max_row+1):
            if sheet.cell(row,4).value==9901:unknown=row;break
        self.assertIsNotNone(unknown)
        self.assertEqual(sheet.cell(unknown,11).value,-1)
        self.assertIn(sheet.cell(unknown,12).value,(None,''))
        self.assertIn(sheet.cell(unknown,13).value,(None,''))
        for sheet in self.values:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type,'e',(sheet.title,cell.coordinate,cell.value))
        for sheet in self.formulas:
            for row in sheet:
                for cell in row:
                    if cell.data_type=='f':
                        self.assertNotIn('#REF!',cell.value)
                        self.assertNotIn('[1]',cell.value)

    def test_same_day_buy_before_sale_is_a_deterministic_daily_formula_order(self):
        sheet=self.values['取引']
        rows=[row for row in range(2,sheet.max_row+1)
              if sheet.cell(row,4).value==9902 and sheet.cell(row,1).value.date()==date(2026,10,2)]
        self.assertEqual([sheet.cell(row,6).value for row in rows],['S','B'])
        sold,bought=rows
        self.assertEqual([sheet.cell(sold,column).value for column in (11,12,13)],[3,101,19])
        self.assertEqual([sheet.cell(bought,column).value for column in (11,12,13)],[4,101,0])

    def test_fund_per_ten_thousand_quote_preserves_units_and_gain(self):
        sheet=self.values['取引'];found={}
        for row in range(2,sheet.max_row+1):
            if sheet.cell(row,4).value==9903:
                found[sheet.cell(row,6).value]=[sheet.cell(row,column).value for column in (7,8,9,11,12,13)]
        self.assertEqual(found['B'],[20000,10002,20004,20000,10002,0])
        self.assertEqual(found['S'],[10000,10500,10500,10000,10002,498])

    def test_cash_nisa_receipts_and_totals_are_actual_cached_calculations(self):
        trade=self.values['取引'];dividend=self.values['配当']
        self.assertEqual(trade['O1'].value.date(),date(2026,10,29))
        self.assertEqual([trade.cell(2,column).value for column in range(16,20)],
                         [1000,2000,10000,123.45])
        self.assertEqual(trade['R3'].value,0);self.assertEqual(trade['S3'].value,0)
        self.assertIn(trade['S5'].value,(None,''))
        self.assertTrue(all(trade.cell(6,column).value in (None,'') for column in range(16,20)))
        kinds=[]
        for row in range(2,7):
            kind=dividend.cell(row,6).value;kinds.append(kind)
            if kind in ('税還','利息'):
                for column in (3,4,5,7,8):self.assertIn(dividend.cell(row,column).value,(None,''))
        self.assertEqual(kinds.count('税還'),2);self.assertEqual(kinds.count('利息'),2)
        blocks=self.model['sheets'][1]['blocks'];total=blocks['cumulative_total_row']
        self.assertEqual([dividend.cell(total,column).value for column in (12,13,14)],[15,27,7])
        self.assertEqual(dividend.cell(blocks['cumulative_last_row'],12).value,49)
        for column in (12,13,14):self.assertTrue(self.formulas['配当'].cell(total,column).value.startswith('=SUM('))
        self.assertEqual(dividend['X2'].value,150);self.assertEqual(dividend['Y2'].value,155)
        label=blocks['calendar_label_row']
        self.assertEqual(dividend.cell(label,11).value,'カレンダー')
        dates=[dividend.cell(row,11).value.date() for row in range(label+1,blocks['calendar_last_row']+1)]
        self.assertEqual(dates,sorted(dates))
        actual=next(row for row in range(label+1,blocks['calendar_last_row']+1)
                    if dividend.cell(row,11).value.date()==date(2026,10,15))
        predicted=next(row for row in range(label+1,blocks['calendar_last_row']+1)
                       if dividend.cell(row,11).value.date()==date(2026,11,10))
        self.assertEqual(dividend.cell(actual,14).value,15)
        self.assertEqual(dividend.cell(predicted,14).value,20)
        self.assertEqual(blocks['calendar_highlight_row'],predicted)


class ActualOfficeLiteralAccountCriteria(unittest.TestCase):
    """A separate real-engine fixture exercises valid wildcard-bearing names."""
    @classmethod
    def setUpClass(cls):
        if load_workbook is None:raise unittest.SkipTest('openpyxl is needed only for optional Office validation')
        executable=office_executable()
        if executable is None:raise unittest.SkipTest('LibreOffice is not installed; set SBI_TEST_OFFICE to its executable')
        temporary=tempfile.TemporaryDirectory(prefix='sbi_v36_literal_accounts_')
        cls.addClassCleanup(temporary.cleanup);directory=Path(temporary.name)
        store=Store(directory/'data');cls.addClassCleanup(store.close)
        activities=Activities(store)
        samples={
            'A':('Account*',2,200,1,150,1,100,50),
            'B':('Account Other',7,2100,2,700,5,300,100),
            'C':('Account?',3,600,1,250,2,200,50),
            'D':('AccountX',11,4400,3,1300,8,400,100),
            'E':('Account~*',5,2500,2,1100,3,500,100),
        }
        cls.expected={};cls.quantities={}
        for account,(name,bought,cost,sold,receipt,left,average,gain) in samples.items():
            store.rename(account,name);store.set_owner(account,'Example '+account)
            rows=[]
            for day,side,quantity,cash in (('2026-10-01','買',bought,cost),
                                            ('2026-10-02','売',sold,receipt)):
                rows.append(dict(trade_date=day,settlement_date=day,name='Literal Criteria Security',
                    code='9904',market='東証',side=side,custody='特定',tax_class='',
                    quantity=str(quantity),price=str(cash/quantity),fees='0',tax='0',
                    settlement_amount=str(cash),asset_type='国内株式',quantity_unit='株',
                    price_unit='円/株',currency='JPY'))
            raw=json.dumps(rows,sort_keys=True).encode()
            store.import_snapshot(account,Snapshot(raw,'literal-'+account+'.csv',
                hashlib.sha256(raw).hexdigest(),'2026-10-01','2026-10-02',rows))
            cls.expected[(name,'B')]=(bought,average,0)
            cls.expected[(name,'S')]=(left,average,gain)
        before='\n'.join(store.db.iterdump())
        model=build(store,activities,'2026-10',today=date(2026,10,29))
        if before!='\n'.join(store.db.iterdump()):raise AssertionError('Literal-name report changed the source ledger')
        original=write_model(model,directory/'source');converted=directory/'recalculated';converted.mkdir()
        result=subprocess.run([executable,'-env:UserInstallation='+(directory/'office-profile').as_uri(),
            '--headless','--convert-to','xlsx','--outdir',str(converted),str(original)],
            stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,timeout=90)
        recalculated=converted/original.name
        if result.returncode!=0 or not recalculated.is_file():
            raise AssertionError('LibreOffice did not recalculate the literal-name fixture: '+result.stdout+result.stderr)
        cls.values=load_workbook(recalculated,data_only=True)
        cls.formulas=load_workbook(recalculated,data_only=False)
        cls.addClassCleanup(cls.values.close);cls.addClassCleanup(cls.formulas.close)

    def test_wildcard_question_mark_and_tilde_names_keep_separate_running_quantities_and_costs(self):
        sheet=self.values['取引'];formulas=self.formulas['取引'];found={}
        for row in range(2,sheet.max_row+1):
            if not isinstance(sheet.cell(row,1).value,datetime):continue
            key=(sheet.cell(row,2).value,sheet.cell(row,6).value)
            found[key]=tuple(sheet.cell(row,column).value for column in (11,12,13))
            for column in (11,12,13):self.assertEqual(formulas.cell(row,column).data_type,'f')
            for cell in sheet[row]:self.assertNotEqual(cell.data_type,'e',(cell.coordinate,cell.value))
        self.assertEqual(found,self.expected)
        self.assertEqual(len(found),10)
